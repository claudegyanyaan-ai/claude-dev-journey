"""Round-trip tests for store.py against real (temp-file) .xlsx workbooks.

Unlike the other test files, this one does real file I/O -- there's no
meaningful way to test "does openpyxl read back what it wrote" as pure
logic. Every test uses its own tempfile.TemporaryDirectory(), so nothing
here ever touches the real data/enrollments.xlsx.
"""

import os
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))
import store  # noqa: E402
import guardrail  # noqa: E402


def make_lead(**overrides):
    lead = {
        "name": "Priya Sharma",
        "age": "28",
        "course_name": "Advanced Python for Data Science",
        "contact": "98765-43210",
        "email": "priya.sharma88@gmail.com",
    }
    lead.update(overrides)
    return lead


class TestLoadRows(unittest.TestCase):
    def test_missing_file_returns_empty_list(self):
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "does_not_exist.xlsx")
            self.assertEqual(store.load_rows(path), [])


class TestAppendAndLoadRoundTrip(unittest.TestCase):
    def test_round_trip_preserves_values(self):
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "enrollments.xlsx")
            store.append_row(make_lead(), path=path)
            rows = store.load_rows(path)
            self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0]["email"], "priya.sharma88@gmail.com")
            self.assertEqual(rows[0]["name"], "Priya Sharma")

    def test_multiple_appends_accumulate_in_order(self):
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "enrollments.xlsx")
            store.append_row(make_lead(name="A", email="a@example.com"), path=path)
            store.append_row(make_lead(name="B", email="b@example.com"), path=path)
            rows = store.load_rows(path)
            self.assertEqual([r["name"] for r in rows], ["A", "B"])

    def test_columns_written_in_fields_order_regardless_of_dict_key_order(self):
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "enrollments.xlsx")
            lead = {"email": "z@example.com", "name": "Z", "age": "40",
                     "contact": "9", "course_name": "Q"}
            store.append_row(lead, path=path)
            rows = store.load_rows(path)
            self.assertEqual(rows[0]["name"], "Z")
            self.assertEqual(rows[0]["email"], "z@example.com")

    def test_missing_field_in_lead_becomes_empty_string_not_a_crash(self):
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "enrollments.xlsx")
            lead = {"name": "Partial", "email": "p@example.com"}  # no age/course_name/contact
            store.append_row(lead, path=path)
            rows = store.load_rows(path)
            self.assertEqual(rows[0]["age"], "")

    def test_creates_parent_directory_if_missing(self):
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "nested", "dir", "enrollments.xlsx")
            store.append_row(make_lead(), path=path)
            self.assertTrue(os.path.isfile(path))

    def test_header_row_matches_guardrail_fields(self):
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "enrollments.xlsx")
            store.append_row(make_lead(), path=path)
            from openpyxl import load_workbook
            wb = load_workbook(path)
            header = [c.value for c in next(wb[store.SHEET_NAME].iter_rows(min_row=1, max_row=1))]
            self.assertEqual(header, guardrail.FIELDS)


class TestIntegrationWithGuardrail(unittest.TestCase):
    def test_loaded_rows_feed_directly_into_duplicate_check(self):
        # store.py's output shape must satisfy what guardrail.py expects,
        # since that's the whole reason existing_rows is a plain argument
        # rather than something guardrail.py reads itself.
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "enrollments.xlsx")
            store.append_row(make_lead(), path=path)
            rows = store.load_rows(path)
            dup = guardrail.find_duplicate_email("priya.sharma88@gmail.com", rows)
            self.assertIsNotNone(dup)
            self.assertEqual(dup["name"], "Priya Sharma")

    def test_no_duplicate_on_empty_spreadsheet(self):
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "does_not_exist_yet.xlsx")
            rows = store.load_rows(path)
            result = guardrail.evaluate(make_lead(), existing_rows=rows)
            self.assertEqual(result["decision"], "auto_write")


if __name__ == "__main__":
    unittest.main()
