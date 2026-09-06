"""Safe Excel read/write for the enrollment spreadsheet.

Only two operations: read every existing row (for guardrail.py's duplicate
check) and append one new row. "Safe" specifically means "can't corrupt or
overwrite data that's already there" -- there is no update or delete here at
all, by design.

File location and column schema weren't pinned down anywhere in the journey
checkpoint or the root planning docs, so this makes a reasonable choice
(data/enrollments.xlsx, one "Enrollments" sheet, columns in the same order
as guardrail.FIELDS) rather than blocking on it -- both are just constants
below, easy to change if a real spreadsheet needs to live somewhere else.
"""

import os
from openpyxl import Workbook, load_workbook

from guardrail import FIELDS  # one source of truth for column order

DEFAULT_PATH = os.path.join(os.path.dirname(__file__), "..", "data", "enrollments.xlsx")
SHEET_NAME = "Enrollments"


def _open_or_create(path):
    """Load the workbook at path, or create a fresh one with a header row."""
    if os.path.isfile(path):
        wb = load_workbook(path)
        if SHEET_NAME not in wb.sheetnames:
            raise ValueError(f"{path} exists but has no '{SHEET_NAME}' sheet.")
        return wb
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.append(FIELDS)
    return wb


def load_rows(path=DEFAULT_PATH):
    """Return every existing enrollment as a list of dicts, in FIELDS order.

    Returns [] if the file doesn't exist yet -- an empty spreadsheet has no
    duplicates to find, which is exactly what guardrail.py's duplicate check
    needs for the very first enrollment ever written.
    """
    if not os.path.isfile(path):
        return []

    wb = load_workbook(path, read_only=True)
    try:
        if SHEET_NAME not in wb.sheetnames:
            raise ValueError(f"{path} exists but has no '{SHEET_NAME}' sheet.")
        ws = wb[SHEET_NAME]

        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []

        header = list(rows[0])
        return [
            {header[i]: ("" if row[i] is None else str(row[i])) for i in range(len(header))}
            for row in rows[1:]
        ]
    finally:
        # read_only=True keeps the underlying zip file handle open until
        # closed explicitly -- on Windows (not Linux) that leaves the file
        # locked, which surfaced as real temp-dir cleanup failures in
        # testing. Always close it, on every exit path.
        wb.close()


def append_row(lead, path=DEFAULT_PATH):
    """Append one lead dict as a new row, creating the file+header if needed.

    Writes values in FIELDS order regardless of the lead dict's own key
    order (or missing keys, which become ""), so the spreadsheet's columns
    stay stable across every call no matter what produced the dict.
    """
    parent = os.path.dirname(path)
    if parent:
        os.makedirs(parent, exist_ok=True)
    wb = _open_or_create(path)
    try:
        ws = wb[SHEET_NAME]
        ws.append([lead.get(field, "") or "" for field in FIELDS])
        wb.save(path)
    finally:
        wb.close()
